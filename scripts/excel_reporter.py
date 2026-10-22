"""Input tables and result workbooks; source rows/sheets are retained as values."""
from dataclasses import dataclass
from pathlib import Path
import re
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .input_parser import extract


@dataclass
class Row:
    sheet: str
    index: int
    value: str
    role: str = "publication"
    input_error: str = ""
    exclusion: str = ""


HEADERS = {"특허번호", "공보번호", "공개번호", "등록번호", "출원번호", "patentno", "patentnumber", "publicationno", "publicationnumber", "applicationno", "applicationnumber"}


def header_key(value):
    return re.sub(r"[\s._()]+", "", str(value or "")).lower()


def read_table(path, column=None, sheet=None, country=None):
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        wb = load_workbook(path, data_only=False)
        sheets = {ws.title: list(ws.values) for ws in wb.worksheets}
    elif path.suffix.lower() == ".xls":
        import pandas as pd
        frames = pd.read_excel(path, sheet_name=None, header=None, dtype=str, keep_default_na=False, engine="xlrd")
        sheets = {name: list(frame.itertuples(index=False, name=None)) for name, frame in frames.items()}
    elif path.suffix.lower() == ".csv":
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp949")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        sheets = {"CSV": list(csv.reader(content.splitlines(), dialect))}
    else:
        raise ValueError("지원 형식: .xlsx, .xls, .csv")
    if sheet and sheet not in sheets:
        raise ValueError("시트를 찾을 수 없습니다: " + sheet)
    rows, layouts = [], {}
    for name, values in sheets.items():
        width = max((len(v) for v in values), default=0)
        chosen = None
        for i, line in enumerate(values[:30]):
            hits = [j for j, v in enumerate(line) if (header_key(v) == header_key(column) if column else header_key(v) in HEADERS)]
            if hits:
                # Publication headers take precedence over application headers.
                hits.sort(key=lambda j: "출원" in str(line[j]) or "application" in str(line[j]).lower())
                chosen = (i, hits[0], str(line[hits[0]]))
                break
        if chosen is None and not column:
            for j in range(width):
                hit = next((i for i, line in enumerate(values) if j < len(line) and extract(line[j] or "", country=country)), None)
                if hit is not None:
                    chosen = (-1, j, "")
                    break
        h, c, label = chosen if chosen else (-1, None, "")
        layouts[name] = (h, width)
        role = "application" if "출원" in label or "application" in label.lower() else "publication"
        for i, line in enumerate(values, 1):
            if not any(v is not None and str(v).strip() for v in line):
                continue
            # Only a recognized header is omitted; never infer a header from a patent hit.
            if i == h + 1:
                continue
            value = (line[c] if c < len(line) else "") if c is not None else " | ".join(str(v) for v in line if v is not None)
            exclusion = (f"--sheet {sheet} 지정으로 대상 외 시트" if sheet and name != sheet else
                         "확인된 번호 헤더 앞의 안내행" if i < h + 1 else "")
            error = (f"지정한 컬럼을 찾을 수 없습니다: {column}" if column else
                     "특허번호 컬럼을 찾지 못했습니다. --column / --country를 확인하세요.") if c is None else ""
            rows.append(Row(name, i, str(value if value is not None else ""), role, error, exclusion))
    if not rows:
        raise ValueError("특허번호 컬럼을 찾지 못했습니다. --column / --country를 지정하세요.")
    return sheets, layouts, rows


def safe_cell(cell, value):
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"  # no execution of user-supplied spreadsheet formulas


def result_values(result):
    labels = dict(official='공식 영문자료', original_only='원문 단독', ai_translation='AI 번역',
                  ai_summary='AI 요약', site_machine_translation='사이트 기계번역')
    return [result['status'], result.get('filename', ''), result.get('pages', ''), result.get('note', ''),
            labels.get(result.get('material_type'), ''), result.get('review_status', ''), result.get('review_packet', ''),
            result.get('reviewed_at', '')]


def write_report(path, sheets, layouts, results):
    wb = Workbook()
    wb.remove(wb.active)
    headers = ["처리결과(생성/검수/대기/실패)", "생성된 PDF 파일명", "총 페이지 수", "비고(실패 사유 또는 출처 사이트)",
               "자료 유형", "최종 검수 상태", "검수 근거 JSON", "검수 일시"]
    for name, values in sheets.items():
        ws = wb.create_sheet(name)
        h, width = layouts.get(name, (0, max((len(v) for v in values), default=0)))
        shift = 1 if h < 0 else 0
        for i, line in enumerate(values, 1 + shift):
            for j, value in enumerate(line, 1):
                safe_cell(ws.cell(i, j), value)
        if name not in layouts:
            continue
        hr = max(1, h + 1)
        for j, label in enumerate(headers, width + 1):
            cell = ws.cell(hr, j, label)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="24456A")
        for result in results:
            if result["sheet"] != name:
                continue
            vals = result_values(result)
            for j, val in enumerate(vals, width + 1):
                safe_cell(ws.cell(result["row"] + shift, j), val)
                ws.cell(result["row"] + shift, j).alignment = Alignment(vertical="top", wrap_text=True)
            if result.get("output_path"):
                cell = ws.cell(result["row"] + shift, width + 2)
                cell.hyperlink = Path(result["output_path"]).resolve().as_uri()
                cell.style = "Hyperlink"
        from openpyxl.utils import get_column_letter
        for offset, size in enumerate([24, 42, 16, 85, 24, 22, 60, 28], width + 1):
            ws.column_dimensions[get_column_letter(offset)].width = size
        ws.freeze_panes = f"A{hr + 1}"
    wb.save(path)


def update_review_report(source, target, results):
    """Write a separate reviewed report; retain the original input and pending report."""
    wb = load_workbook(source)
    for result in results:
        ws = wb[result['sheet']]
        for j, value in enumerate(result_values(result), result['report_start_column']):
            safe_cell(ws.cell(result['report_row'], j), value)
    wb.save(target)
