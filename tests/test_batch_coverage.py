"""Every nonblank input row must have an outcome, apart from recognized headers."""
from contextlib import contextmanager, redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from scripts.build_references import parser, run
from scripts.excel_reporter import read_table


class BatchCoverage(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def table(self, suffix, sheets):
        path = self.root / ('input' + suffix)
        if suffix == '.csv':
            import csv
            with path.open('w', encoding='utf-8-sig', newline='') as f:
                csv.writer(f).writerows(next(iter(sheets.values())))
        elif suffix == '.xlsx':
            wb = Workbook(); wb.remove(wb.active)
            for name, lines in sheets.items():
                ws = wb.create_sheet(name)
                for line in lines: ws.append(line)
            wb.save(path)
        else:
            import xlwt
            wb = xlwt.Workbook()
            for name, lines in sheets.items():
                ws = wb.add_sheet(name)
                for i, line in enumerate(lines):
                    for j, value in enumerate(line):
                        if value is not None: ws.write(i, j, value)
            wb.save(str(path))
        return path

    def execute(self, path, *extra):
        out = self.root / 'result'
        args = parser().parse_args(['--input', str(path), '--out', str(out), *extra])
        # A deterministic startup failure substitutes for network collection here.
        @contextmanager
        def offline(_):
            raise RuntimeError('fixture offline')
            yield  # context fails on entry, like the actual browser_session
        with patch('scripts.build_references.browser_session', side_effect=offline) as browser:
            with redirect_stdout(io.StringIO()) as stdout:
                code = run(args)
        return out, code, stdout.getvalue(), browser.call_count

    def test_headerless_leading_invalid_rows_all_formats(self):
        for suffix in ('.csv', '.xlsx', '.xls'):
            with self.subTest(suffix=suffix):
                path = self.table(suffix, {'data': [('bad',), (None,), ('JP2006198019A',), ('', 'memo')]})
                _, layouts, rows = read_table(path)
                self.assertEqual([r.index for r in rows], [1, 3, 4])
                self.assertEqual(rows[0].value, 'bad')
                self.assertEqual(next(iter(layouts.values()))[0], -1)

    def test_dry_run_keeps_leading_failure_and_invalid_sheet(self):
        path = self.table('.xlsx', {'data': [('bad',), ('JP2006198019A',)], 'unknown': [('unrecognized',), ('still bad',)]})
        out, code, _, calls = self.execute(path, '--dry-run')
        self.assertEqual((code, calls), (1, 0))
        plan = json.loads((out / 'input-plan.json').read_text(encoding='utf-8'))
        self.assertEqual([(r['sheet'], r['row']) for r in plan], [('data', 1), ('data', 2), ('unknown', 1), ('unknown', 2)])
        self.assertEqual([r['status'] for r in plan], ['실패', '처리 예정', '실패', '실패'])
        counts = json.loads((out / 'input-accounting.json').read_text(encoding='utf-8'))
        self.assertEqual(counts['input_rows'], 4)
        self.assertEqual(counts['remaining_rows'], 0)

    def test_report_and_summary_cover_mixed_sheets(self):
        path = self.table('.xlsx', {'data': [('bad',), ('JP2006198019A',)], 'unknown': [('unrecognized',), (0,)]})
        out, code, stdout, calls = self.execute(path, '--zip')
        self.assertEqual((code, calls), (1, 1))
        rows = json.loads(next(out.glob('처리결과_*.json')).read_text(encoding='utf-8'))
        self.assertEqual(len(rows), 4)
        self.assertTrue(all(r['status'] == '실패' and r['note'] for r in rows))
        wb = load_workbook(next(out.glob('처리결과_*.xlsx')))
        self.assertEqual(wb['data']['A2'].value, 'bad')
        self.assertEqual(wb['data']['B2'].value, '실패')
        self.assertEqual(wb['unknown']['A3'].value, 0)
        self.assertEqual(wb['unknown']['B3'].value, '실패')
        summary = json.loads(stdout.splitlines()[-1])
        self.assertEqual((summary['input_rows'], summary['recorded_rows'], summary['failed']), (4, 4, 4))
        counts = json.loads(Path(summary['accounting']).read_text(encoding='utf-8'))
        self.assertEqual(counts['statuses'], {'실패': 4})
        import zipfile
        with zipfile.ZipFile(next(out.glob('*.zip'))) as z:
            self.assertEqual(len([n for n in z.namelist() if n.endswith('.json')]), 1)

    def test_all_unrecognized_rows_produce_report_without_browser(self):
        path = self.table('.csv', {'data': [('bad',), ('also bad',)]})
        out, code, _, calls = self.execute(path)
        self.assertEqual((code, calls), (1, 0))
        rows = json.loads(next(out.glob('처리결과_*.json')).read_text(encoding='utf-8'))
        self.assertEqual([r['row'] for r in rows], [1, 2])

    def test_recognized_header_preamble_and_selected_out_sheet(self):
        path = self.table('.xlsx', {'selected': [('title',), ('Patent No.',), ('bad',)],
                                    'other': [('Publication No.',), ('JP2006198019A',)]})
        out, code, stdout, calls = self.execute(path, '--sheet', 'selected')
        self.assertEqual((code, calls), (1, 0))
        rows = json.loads(next(out.glob('처리결과_*.json')).read_text(encoding='utf-8'))
        self.assertEqual([(r['sheet'], r['row'], r['status']) for r in rows],
                         [('selected', 1, '제외'), ('selected', 3, '실패'), ('other', 2, '제외')])
        self.assertIn('--sheet', rows[-1]['note'])
        self.assertEqual(json.loads(stdout.splitlines()[-1])['excluded'], 2)
        wb = load_workbook(next(out.glob('처리결과_*.xlsx')))
        self.assertEqual(wb['selected']['B1'].value, '제외')
        self.assertEqual(wb['other']['B2'].value, '제외')

    def test_missing_explicit_column_recorded_per_sheet(self):
        path = self.table('.xlsx', {'good': [('Patent No.',), ('bad',)], 'other': [('memo',), ('JP2006198019A',)]})
        out, code, _, calls = self.execute(path, '--column', 'Patent No.')
        self.assertEqual((code, calls), (1, 0))
        rows = json.loads(next(out.glob('처리결과_*.json')).read_text(encoding='utf-8'))
        self.assertEqual(len(rows), 3)
        self.assertIn('지정한 컬럼', rows[-1]['note'])


if __name__ == '__main__': unittest.main()
