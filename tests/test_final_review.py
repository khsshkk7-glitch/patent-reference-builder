"""Review state is explicit and bound to the final PDF, source parts and rendered evidence."""
from contextlib import nullcontext, redirect_stdout
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from openpyxl import load_workbook
from scripts.build_references import parser, run
from scripts.collector_common import Collected
from scripts.final_review import CHECKS, apply_review, render_packet, write_json
from scripts.input_parser import one
from scripts.pdf_merger import sha256
from tests.test_core import pdf


class FinalReview(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.b = self.root / 'front.pdf'; self.o = self.root / 'original.pdf'
        pdf(self.b, [595.28]); pdf(self.o, [400, 400, 400, 400])
        entry = dict(publication='JP2006198019A', source='Espacenet',
            biblio_pdf=str(self.b), original_pdf=str(self.o), original_pages=4,
            biblio_url='https://worldwide.espacenet.com/publicationDetails/biblio',
            original_url='https://worldwide.espacenet.com/publicationDetails/originalDocument',
            biblio_sha256=sha256(self.b), original_sha256=sha256(self.o), identity_verified=True,
            english_abstract_verified=True, full_document_verified=True, layout_verified=True,
            verification_note='Synthetic parts reviewed, not a final merged PDF review')
        manifest = self.root / 'manifest.json'; write_json(manifest, {'JP2006198019A':entry})
        self.out = self.root / 'out'
        args = parser().parse_args(['--text', 'JP2006198019A,JP2006198019A,bad', '--manifest', str(manifest),
                                   '--out', str(self.out), '--zip'])
        with redirect_stdout(io.StringIO()) as stdout: self.assertEqual(run(args), 1)
        self.summary = json.loads(stdout.getvalue().splitlines()[-1])
        self.results = next(self.out.glob('처리결과_*.json'))
        self.rows = json.loads(self.results.read_text(encoding='utf-8'))
        self.packet = Path(self.rows[0]['review_packet'])

    def draft(self):
        packet = render_packet(self.packet)
        packet.update(decision='approved', reviewer='Synthetic QA reviewer', review_note='Fixture review, not real patent certification')
        packet['checks'] = dict.fromkeys(CHECKS, True)
        packet['sections'] = dict(claims=dict(present=True, pages=[3], note='Claim boundary fixture on final page 3'),
                                  drawings=dict(present=True, pages=[5], note='Drawing boundary fixture on final page 5'))
        write_json(self.packet, packet)
        return packet

    def approved(self):
        packet = self.draft()
        render_packet(self.packet, [3])
        packet = json.loads(self.packet.read_text(encoding='utf-8'))
        packet.update(decision='approved', reviewer='Synthetic QA reviewer', review_note='Fixture review only')
        packet['checks'] = dict.fromkeys(CHECKS, True)
        write_json(self.packet, packet)
        return packet

    def test_generation_stays_pending_even_with_reviewed_input_manifest(self):
        self.assertEqual([r['status'] for r in self.rows], ['생성 완료', '생성 완료', '실패'])
        self.assertTrue(self.rows[0]['official_materials_complete'])
        self.assertFalse(self.rows[0]['complete_reference'])
        self.assertFalse(self.rows[0]['final_review_complete'])
        self.assertEqual(self.summary['success'], 0)
        self.assertEqual(self.summary['final_review_pending'], 2)
        self.assertEqual(self.rows[0]['review_packet'], self.rows[1]['review_packet'])
        packet = json.loads(self.packet.read_text(encoding='utf-8'))
        self.assertFalse(any(packet['checks'].values()))
        self.assertEqual(packet['images'], [])
        wb = load_workbook(self.results.with_suffix('.xlsx'))
        self.assertEqual(wb.active['G2'].value, '검수 대기')
        with zipfile.ZipFile(next(self.out.glob('검수대기_*.zip'))) as z:
            self.assertIn('검수안내.txt', z.namelist())
            self.assertIn('검수', z.read('검수안내.txt').decode('utf-8'))
        self.assertFalse(list(self.out.glob('검수완료*.zip')))

    def test_online_and_local_use_same_final_review_gate(self):
        parts = Collected(one('JP2006198019A'), self.b, self.o, 4, 'Espacenet',
                          'https://worldwide.espacenet.com/biblio', 'https://worldwide.espacenet.com/original',
                          'Final visual review required')
        out = self.root / 'online'
        args = parser().parse_args(['--text','JP2006198019A','--out',str(out),'--delay','0'])
        with patch('scripts.build_references.browser_session', return_value=nullcontext(object())), \
             patch('scripts.build_references.collect', return_value=parts), redirect_stdout(io.StringIO()):
            self.assertEqual(run(args), 0)  # generation exit code, not review approval
        record = json.loads(next(out.glob('*.json')).read_text(encoding='utf-8'))[0]
        for key in ('review_status', 'final_review_complete', 'complete_reference', 'material_type'):
            self.assertEqual(record[key], self.rows[0][key])

    def test_approval_requires_boundary_images_and_all_checks(self):
        packet = self.draft()
        with self.assertRaisesRegex(ValueError, '이미지 누락'): apply_review(self.results, self.packet)
        packet = self.approved()
        for key in CHECKS:
            bad = dict(packet, checks=dict(packet['checks'], **{key:False}))
            write_json(self.packet, bad)
            with self.assertRaisesRegex(ValueError, '검수 미완료'): apply_review(self.results, self.packet)
        self.assertFalse(list(self.out.glob('검수완료*.zip')))

    def test_approved_report_and_zip_preserve_hashes_and_duplicate_rows(self):
        self.approved()
        result = apply_review(self.results, self.packet)
        rows = json.loads(Path(result['results']).read_text(encoding='utf-8'))
        self.assertEqual([r['status'] for r in rows], ['검수 완료', '검수 완료', '실패'])
        self.assertTrue(rows[0]['complete_reference'])
        self.assertEqual(rows[0]['review_evidence']['original_sha256'], sha256(self.o))
        self.assertTrue(rows[0]['reviewed_at'])
        self.assertEqual(load_workbook(result['report']).active['G2'].value, '검수 완료')
        with zipfile.ZipFile(result['archive']) as z:
            self.assertEqual(sum(n.endswith('.pdf') for n in z.namelist()), 1)
            self.assertEqual(sum(n.endswith('.png') for n in z.namelist()), 4)
            self.assertTrue(any(n.endswith('/검수기록.json') for n in z.namelist()))
        # Original pending report remains an accurate snapshot of the generation stage.
        self.assertFalse(json.loads(self.results.read_text(encoding='utf-8'))[0]['final_review_complete'])

    def test_rejection_for_name_identity_clipping_and_missing_drawings(self):
        for note in ('Inventor name misread', 'Publication number misread', 'Front page clipped', 'Original drawings missing'):
            with self.subTest(note=note):
                packet = json.loads(self.packet.read_text(encoding='utf-8'))
                packet.update(decision='rejected', reviewer='Synthetic reviewer', review_note=note)
                write_json(self.packet, packet)
                result = apply_review(self.results, self.packet)
                rows = json.loads(Path(result['results']).read_text(encoding='utf-8'))
                self.assertEqual(rows[0]['review_status'], '검수 반려')
                self.assertFalse(rows[0]['complete_reference']); self.assertIsNone(result['archive'])
                self.assertIn(note, rows[0]['note'])

    def test_changed_pdf_or_source_is_not_approved(self):
        self.approved()
        for target in (Path(self.rows[0]['output_path']), self.b, self.o):
            with self.subTest(target=target):
                original = target.read_bytes(); target.write_bytes(original + b'changed')
                with self.assertRaisesRegex(ValueError, '해시 불일치'): apply_review(self.results, self.packet)
                target.write_bytes(original)

    def test_changed_images_and_forged_bindings_are_not_approved(self):
        packet = self.approved()
        image = Path(packet['images'][0]['path']); original = image.read_bytes()
        image.write_bytes(b'not a PNG')
        with self.assertRaisesRegex(ValueError, '이미지 변경'): apply_review(self.results, self.packet)
        image.write_bytes(original)
        packet['material_type'] = 'ai_translation'; write_json(self.packet, packet)
        with self.assertRaisesRegex(ValueError, '바인딩 불일치'): apply_review(self.results, self.packet)

    def test_missing_sections_notes_or_reviewer_block_approval(self):
        packet = self.approved()
        for key, value in (('reviewer', ''), ('review_note', ''), ('decision', 'pending')):
            write_json(self.packet, dict(packet, **{key:value}))
            with self.assertRaises(ValueError): apply_review(self.results, self.packet)
        packet['sections']['drawings'] = dict(present=False, pages=[], note='')
        write_json(self.packet, packet)
        with self.assertRaisesRegex(ValueError, '부재 근거'): apply_review(self.results, self.packet)

    def test_rendering_does_not_approve_and_rerender_invalidates_decision(self):
        packet = self.approved()
        packet = render_packet(self.packet)
        self.assertEqual(packet['decision'], 'pending')
        self.assertFalse(any(packet['checks'].values()))
        with self.assertRaises(ValueError): apply_review(self.results, self.packet)

    def test_ai_and_site_translation_remain_assisted_after_review(self):
        # Classification mechanics only; these blank fixtures are not actual translated documents.
        for material in ('ai_translation', 'ai_summary', 'site_machine_translation', 'original_only'):
            with self.subTest(material=material):
                rows = json.loads(self.results.read_text(encoding='utf-8'))
                for record in rows[:2]:
                    record.update(material_type=material, official_materials_complete=False)
                write_json(self.results, rows)
                packet = self.approved(); packet['material_type'] = material
                write_json(self.packet, packet)
                result = apply_review(self.results, self.packet)
                reviewed = json.loads(Path(result['results']).read_text(encoding='utf-8'))[0]
                self.assertTrue(reviewed['final_review_complete'])
                self.assertFalse(reviewed['complete_reference'])
                self.assertEqual(reviewed['material_type'], material)

    def test_prior_approved_file_change_blocks_cumulative_zip(self):
        manifest = json.loads((self.root / 'manifest.json').read_text(encoding='utf-8'))
        manifest['EP1234567B1'] = dict(manifest['JP2006198019A'], publication='EP1234567B1')
        write_json(self.root / 'two.json', manifest)
        out = self.root / 'two'
        args = parser().parse_args(['--text','JP2006198019A,EP1234567B1', '--manifest',str(self.root / 'two.json'), '--out',str(out)])
        with redirect_stdout(io.StringIO()): self.assertEqual(run(args), 0)
        self.results = next(out.glob('*.json'))
        rows = json.loads(self.results.read_text(encoding='utf-8'))
        self.packet = Path(rows[0]['review_packet']); self.approved()
        first = apply_review(self.results, self.packet)
        self.assertEqual((first['reviewed_files'], first['pending_rows']), (1, 1))
        self.packet = Path(rows[1]['review_packet']); self.approved()
        target = Path(rows[0]['output_path']); target.write_bytes(target.read_bytes() + b'changed')
        with self.assertRaisesRegex(ValueError, '해시 불일치'):
            apply_review(first['results'], self.packet)
        self.assertEqual(len(list(out.glob('검수완료자료_*.zip'))), 1)

    def test_original_only_review_has_no_front_and_preserves_material_type(self):
        out = self.root / 'original-only'
        parts = Collected(one('JP2006198019A'), None, self.o, 4, 'Espacenet', '',
                          'https://worldwide.espacenet.com/original', 'Original-only fixture')
        args = parser().parse_args(['--text','JP2006198019A','--out',str(out),'--original-only','--delay','0'])
        with patch('scripts.build_references.browser_session', return_value=nullcontext(object())), \
             patch('scripts.build_references.espacenet_collector.collect_original_only', return_value=parts), redirect_stdout(io.StringIO()):
            self.assertEqual(run(args), 0)
        results = next(out.glob('*.json')); row = json.loads(results.read_text(encoding='utf-8'))[0]
        packet_path = Path(row['review_packet']); packet = render_packet(packet_path)
        self.assertEqual({i['page'] for i in packet['images']}, {1, 4})
        packet.update(decision='approved', reviewer='Synthetic QA', review_note='Original-only fixture reviewed')
        packet['checks'] = dict.fromkeys(CHECKS, True)
        packet['sections'] = {key:dict(present=False, pages=[], note='No such section in synthetic fixture') for key in ('claims', 'drawings')}
        write_json(packet_path, packet)
        result = apply_review(results, packet_path)
        reviewed = json.loads(Path(result['results']).read_text(encoding='utf-8'))[0]
        self.assertEqual(reviewed['material_type'], 'original_only')
        self.assertTrue(reviewed['final_review_complete']); self.assertFalse(reviewed['complete_reference'])


if __name__ == '__main__': unittest.main()
