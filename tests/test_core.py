import json
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook, load_workbook
from pypdf import PdfReader, PdfWriter
from scripts.input_parser import extract, one, resolve
from scripts.excel_reporter import read_table, write_report
from scripts.pdf_merger import merge, sha256
from scripts.collector_common import from_manifest
from scripts.build_references import parser, run, input_rows


def pdf(path, widths):
    w = PdfWriter()
    for width in widths:
        w.add_blank_page(width=width, height=841.89)
    w.write(path)
    w.close()


class Parsing(unittest.TestCase):
    def test_requested_variants(self):
        for s in ['KR1020200137353A', 'kr 10-2020-0137353 A', 'ＫＲ１０２０２００１３７３５３Ａ']:
            self.assertEqual(one(s).key, 'KR20200137353A')
        self.assertEqual(one('KR 10-2020-0137353').kind, '')
        for s in ['JP 2006-198019 A', '特開2006-198019', 'JP2006198019 (A)']:
            self.assertEqual(one(s).key, 'JP2006198019A')

    def test_countries_and_names(self):
        for s, expected in [('US 2020/0123456 A1','US20200123456A1'), ('EP 1234567 B1','EP1234567B1'),
                            ('WO 2020/123456 A1','WO2020123456A1'), ('CN 112345678 B','CN112345678B')]:
            self.assertEqual(one(s).key, expected)
        self.assertEqual(one('KR1020200137353A').filename, 'KR 10-2020-0137353.pdf')
        self.assertEqual(one('JP 2006-198019 A').filename, 'JP 2006-198019 A.pdf')

    def test_wrong_kind_or_ambiguous_rejected(self):
        with self.assertRaises(ValueError):
            resolve(one('EP1234567A1'), [one('EP1234567B1')])
        with self.assertRaises(ValueError):
            resolve(one('EP1234567'), [one('EP1234567A1'), one('EP1234567B1')])
        with self.assertRaises(ValueError):
            resolve(one('KR1020190063559',role='application'), [one('KR1020200137353A')])
        self.assertEqual(extract('1020200137353'), [])
        self.assertEqual(one('1020200137353', country='KR').key, 'KR20200137353')

    def test_multiple_in_chat_and_invalid(self):
        args = parser().parse_args(['--text','JP2006198019A EP1234567B1, bad','--out','unused'])
        _, _, rows = input_rows(args)
        self.assertEqual([r.index for r in rows], [2,3,4])
        self.assertEqual(rows[-1].value,'bad')

    def test_empty_batch_is_not_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            args=parser().parse_args(['--text',' , ','--out',tmp])
            with self.assertRaises(ValueError): run(args)


class Files(unittest.TestCase):
    def test_human_wait_stops_further_collection_and_reports_unattempted_rows(self):
        from contextlib import nullcontext
        from unittest.mock import patch
        from scripts.browser_session import HumanWaitTimeout
        with tempfile.TemporaryDirectory() as tmp:
            args=parser().parse_args(['--text','JP2004121879A,CN113332661A','--out',tmp])
            with patch('scripts.build_references.browser_session',return_value=nullcontext(object())), \
                 patch('scripts.build_references.collect',side_effect=HumanWaitTimeout('인증 반복')) as collect:
                self.assertEqual(run(args),1)
            self.assertEqual(collect.call_count,1)
            results=json.loads(next(Path(tmp).glob('*.json')).read_text(encoding='utf-8'))
            self.assertEqual(results[0]['note'],'인증 반복')
            self.assertIn('후속 온라인 수집 미실행',results[1]['note'])
            self.assertNotIn('브라우저 시작 실패',results[1]['note'])

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_merge_preserves_order_count_size_and_no_overwrite(self):
        b,o,out = [self.root/n for n in ['b.pdf','o.pdf','out.pdf']]
        pdf(b,[595.28,595.28]); pdf(o,[400,500,600])
        stats=merge(b,o,out,3)
        self.assertEqual(stats['pages'],5)
        self.assertEqual([float(p.mediabox.width) for p in PdfReader(out).pages],[595.28,595.28,400,500,600])
        with self.assertRaises(FileExistsError): merge(b,o,out,3)
        with self.assertRaises(ValueError): merge(b,o,self.root/'wrong.pdf',4)
        self.assertFalse((self.root/'wrong.pdf').exists())

    def test_xlsx_rows_and_failure_mapping(self):
        path=self.root/'input.xlsx'; wb=Workbook(); ws=wb.active; ws.title='목록'
        ws.append(['제목']); ws.append(['담당','공개번호','메모'])
        ws.append(['A','KR1020200137353A','x']); ws.append(['B','invalid','y']);ws.append(['C',None,'번호 없음'])
        wb.create_sheet('참고').append(['보존'])
        wb.save(path)
        sheets,layouts,rows=read_table(path)
        self.assertEqual([(r.sheet,r.index) for r in rows],[('목록',1),('목록',3),('목록',4),('목록',5),('참고',1)])
        self.assertTrue(rows[0].exclusion)
        self.assertTrue(rows[-1].input_error)
        result=self.root/'result.xlsx'
        write_report(result,sheets,layouts,[dict(sheet='목록',row=4,status='실패',note='잘못된 번호')])
        out=load_workbook(result)
        self.assertEqual(out['목록']['D4'].value,'실패')
        self.assertEqual(out['목록']['G4'].value,'잘못된 번호')
        self.assertEqual(out['참고']['A2'].value,'보존')
        self.assertEqual(load_workbook(path)['목록'].max_column,3)

    def test_headerless_csv_and_formula_safety(self):
        p=self.root/'input.csv';p.write_text('JP2006198019A,=HYPERLINK("evil")\nEP1234567B1,x\n',encoding='utf-8-sig')
        sheets,layouts,rows=read_table(p)
        self.assertEqual([r.index for r in rows],[1,2])
        out=self.root/'result.xlsx'
        write_report(out,sheets,layouts,[dict(sheet='CSV',row=1,status='성공',filename='JP.pdf',pages=3,note='Espacenet')])
        w=load_workbook(out)
        self.assertEqual(w.active['A2'].value,'JP2006198019A')
        self.assertEqual(w.active['C2'].value,'성공')
        self.assertEqual(w.active['B2'].data_type,'s')

    def test_application_header_retained(self):
        p=self.root/'app.csv';p.write_text('출원번호\nKR1020190063559\n',encoding='utf-8')
        self.assertEqual(read_table(p)[2][0].role,'application')

    def test_legacy_xls_all_sheets(self):
        import xlwt
        path=self.root/'legacy.xls'; w=xlwt.Workbook()
        for title,value in [('KR','KR1020200137353A'),('JP','JP2006198019A')]:
            s=w.add_sheet(title);s.write(0,0,'Patent No.');s.write(1,0,value)
        w.save(str(path))
        sheets,layouts,rows=read_table(path)
        self.assertEqual([r.sheet for r in rows],['KR','JP'])
        self.assertEqual([r.value for r in rows],['KR1020200137353A','JP2006198019A'])

    def test_manifest_and_batch_duplicate_failure_zip(self):
        b=self.root/'b.pdf';o=self.root/'o.pdf';pdf(b,[595.28]);pdf(o,[595.28,595.28])
        # Synthetic fixtures test mechanics only, not real document identity.
        item=dict(publication='JP2006198019A',source='Espacenet',biblio_pdf='b.pdf',original_pdf='o.pdf',
            biblio_url='https://worldwide.espacenet.com/publicationDetails/biblio',
            original_url='https://worldwide.espacenet.com/publicationDetails/originalDocument',original_pages=2,
            biblio_sha256=sha256(b),original_sha256=sha256(o),identity_verified=True,english_abstract_verified=True,
            full_document_verified=True,layout_verified=True,verification_note='Synthetic fixture, test only')
        manifest=self.root/'review.json';manifest.write_text(json.dumps({'JP2006198019A':item}),encoding='utf-8')
        args=parser().parse_args(['--text','JP2006198019A,JP2006198019A,bad','--manifest',str(manifest),'--out',str(self.root/'out'),'--zip'])
        self.assertEqual(run(args),1)
        result=json.loads(next((self.root/'out').glob('*.json')).read_text(encoding='utf-8'))
        self.assertEqual([r['status'] for r in result],['성공','성공','실패'])
        self.assertEqual(len(list((self.root/'out').glob('*.pdf'))),1)
        self.assertEqual(len(list((self.root/'out').glob('*.zip'))),1)
        item['original_sha256']='invalid'
        from scripts.browser_session import CollectionError
        with self.assertRaises(CollectionError): from_manifest(one('JP2006198019A'),item,self.root)


if __name__=='__main__': unittest.main()
